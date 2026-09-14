import openpyxl

wb_formula = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=False)
wb_value = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=True)

with open('core_tinh_formulas.txt', 'w', encoding='utf-8') as out:
    for sheetname in wb_formula.sheetnames:
        ws_f = wb_formula[sheetname]
        ws_v = wb_value[sheetname]
        out.write(f"=== SHEET: {sheetname} ===\n")
        
        # Header
        headers = [ws_f.cell(1, col).value for col in range(1, ws_f.max_column + 1)]
        out.write(f"Headers: {headers}\n\n")
        
        # All rows
        for r in range(1, ws_f.max_row + 1):
            loai = str(ws_f.cell(r, 1).value).strip()
            field_name = ws_f.cell(r, 2).value if ws_f.max_column >= 2 else ''
            # Check if this row is 'Tính' or has formula in col M or any col
            col_m_f = ws_f.cell(r, 13).value if ws_f.max_column >= 13 else None # Col M is 13? Let's check col letters
            col_l_f = ws_f.cell(r, 12).value if ws_f.max_column >= 12 else None
            
            # Find all formulas in this row
            formulas = {}
            values = {}
            for col in range(1, ws_f.max_column + 1):
                col_letter = openpyxl.utils.get_column_letter(col)
                f = ws_f.cell(r, col).value
                v = ws_v.cell(r, col).value
                if str(f).startswith('='):
                    formulas[col_letter] = f
                    values[col_letter] = v
            
            # If row is marked 'Tính' or has formulas, or just print all rows with their field names
            out.write(f"Row {r:3d} | Loại: {loai:<6} | Chỉ tiêu: {field_name}\n")
            if formulas:
                for c_let, f_str in formulas.items():
                    val = values.get(c_let)
                    out.write(f"     {c_let}{r}: [Formula] {f_str} => [Value] {val}\n")

print("Done extracting formulas")
