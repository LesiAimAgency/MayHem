import openpyxl

wb_f = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=False)
wb_v = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=True)

ws_f = wb_f['Báo cáo từng cty']
ws_v = wb_v['Báo cáo từng cty']

with open('bao_cao_tung_cty.txt', 'w', encoding='utf-8') as out:
    for r in range(1, 60):
        loai = str(ws_f.cell(r, 1).value).strip()
        field = str(ws_f.cell(r, 2).value).strip()
        f_m = str(ws_f.cell(r, 13).value).strip() # Col M (2025)
        v_m = str(ws_v.cell(r, 13).value).strip()
        f_l = str(ws_f.cell(r, 12).value).strip() # Col L (2024)
        v_l = str(ws_v.cell(r, 12).value).strip()
        out.write(f"Row {r:2d} | [{loai:<4}] | {field:<55} | L(2024): {v_l:<15} (fx: {f_l}) | M(2025): {v_m:<15} (fx: {f_m})\n")

print("Done bao_cao_tung_cty")
