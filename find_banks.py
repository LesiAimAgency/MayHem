import openpyxl

wb = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=False)
ws = wb['Báo cáo từng cty']

with open('bank_sections.txt', 'w', encoding='utf-8') as out:
    for r in range(1, ws.max_row + 1):
        c1 = ws.cell(r, 1).value
        c2 = ws.cell(r, 2).value
        if c1 == 'Loại':
            out.write(f"Header at Row {r:3d}: Bank = {c2}\n")

print("Done bank sections")
