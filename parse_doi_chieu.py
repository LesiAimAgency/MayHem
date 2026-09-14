import openpyxl

wb_v = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=True)
ws = wb_v['Đối chiếu dữ liệu']

with open('doi_chieu_data.txt', 'w', encoding='utf-8') as out:
    out.write(f"Max row: {ws.max_row}, max col: {ws.max_column}\n")
    for r in range(1, min(50, ws.max_row + 1)):
        vals = [str(ws.cell(r, col).value) for col in range(1, min(15, ws.max_column + 1))]
        out.write(f"Row {r:2d}: {' | '.join(vals)}\n")

print("Done doi chieu")
