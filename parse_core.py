import openpyxl

wb_formula = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=False)
wb_value = openpyxl.load_workbook(r'c:\MAMP\htdocs\MayHem\backend\core.xlsx', data_only=True)

with open('core_excel_analysis.txt', 'w', encoding='utf-8') as out:
    for sheetname in wb_formula.sheetnames:
        ws_f = wb_formula[sheetname]
        ws_v = wb_value[sheetname]
        out.write(f"=== SHEET: {sheetname} (max_row={ws_f.max_row}, max_col={ws_f.max_column}) ===\n")
        
        # Print header
        header_f = [str(ws_f.cell(1, col).value) for col in range(1, ws_f.max_column + 1)]
        out.write(f"Header: {' | '.join(header_f)}\n\n")
        
        for r in range(1, ws_f.max_row + 1):
            loai = ws_f.cell(r, 1).value
            row_data = []
            for col in range(1, ws_f.max_column + 1):
                f_val = ws_f.cell(r, col).value
                v_val = ws_v.cell(r, col).value
                col_letter = openpyxl.utils.get_column_letter(col)
                if str(f_val).startswith('='):
                    row_data.append(f"{col_letter}{r}: [Formula: {f_val} => Value: {v_val}]")
                else:
                    row_data.append(f"{col_letter}{r}: {f_val}")
            out.write(f"Row {r:2d} ({loai}):\n  " + "\n  ".join(row_data) + "\n\n")

print("Done analyzing core.xlsx")
